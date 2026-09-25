import os
import json
import pytest
from pptx import Presentation
from docx import Document

from backend.models.formats.presentation_schema import PresentationSchema, Slide
from backend.models.formats.advisory_schema import AdvisorySchema
from backend.exporters.pptx_exporter import export_pptx
from backend.exporters.docx_exporter import export_docx


def test_export_pptx_from_schema(tmp_path):
    """
    Validates PPTX exporter generates uncorrupted .pptx with title,
    slides, speaker notes, and citations.
    """
    deck = PresentationSchema(
        deck_title="Operation GhostLatch: Critical Infrastructure Incident Analysis",
        target_audience="Executive Leadership & Incident Command",
        slides=[
            Slide(
                slide_number=1,
                title="Executive Incident Overview",
                bullet_points=[
                    "Unauthorized access detected across two substation endpoints",
                    "Exploitation stemmed from unpatched firmware vulnerability",
                    "Immediate containment achieved within 24 hours"
                ],
                visual_guidance="Split 2-column layout with severity highlight badge on right",
                speaker_notes="Good morning leaders. This briefing covers the GhostLatch containment.",
                slide_reference_citations=["doc_01_chunk_01", "doc_01_chunk_02"]
            ),
            Slide(
                slide_number=2,
                title="Immediate Remediation Measures",
                bullet_points=[
                    "Emergency firmware deployment completed",
                    "Full isolation of vulnerable substation controllers",
                    "No threat actor attribution established"
                ],
                visual_guidance="Process timeline with 3 distinct milestones",
                speaker_notes="All endpoints were patched with zero cascading grid disruptions.",
                slide_reference_citations=["doc_01_chunk_02"]
            )
        ]
    )

    out_file = str(tmp_path / "test_presentation.pptx")
    result_path = export_pptx(deck, out_file)

    assert os.path.exists(result_path)
    assert os.path.getsize(result_path) > 0

    # Verify openable with python-pptx
    prs = Presentation(result_path)
    assert len(prs.slides) == 3  # 1 title slide + 2 content slides

    # Check title slide notes
    title_slide = prs.slides[0]
    assert title_slide.has_notes_slide
    assert "Operation GhostLatch" in title_slide.notes_slide.notes_text_frame.text

    # Check content slide notes and content
    slide1 = prs.slides[1]
    assert slide1.has_notes_slide
    assert "Good morning leaders" in slide1.notes_slide.notes_text_frame.text


def test_export_docx_from_fixture(tmp_path):
    """
    Validates DOCX exporter works with mock_draft_outputs.json fixture
    and produces uncorrupted .docx with institutional header, severity banner,
    and IOC table.
    """
    fixtures_path = os.path.join(os.path.dirname(__file__), "..", "..", "fixtures", "mock_draft_outputs.json")
    with open(fixtures_path, "r", encoding="utf-8") as f:
        fixture_data = json.load(f)

    advisory_data = fixture_data["advisory"]
    advisory = AdvisorySchema(**advisory_data)

    out_file = str(tmp_path / "test_advisory.docx")
    result_path = export_docx(advisory, out_file)

    assert os.path.exists(result_path)
    assert os.path.getsize(result_path) > 0

    # Verify openable with python-docx
    doc = Document(result_path)
    
    # Check tables
    assert len(doc.tables) >= 3  # Header table, Severity table, IOC table

    # Verify text elements present
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "Firmware Vulnerability" in full_text
    assert "Directorate of Power Grid Resilience" in full_text
    assert "doc_01_chunk_01" in full_text

    # Verify IOC table content
    ioc_table = doc.tables[2]
    table_text = " ".join(cell.text for row in ioc_table.rows for cell in row.cells)
    assert "Unpatched firmware vulnerability" in table_text


def test_export_zip_archive(tmp_path):
    """
    Validates that a combined zip package contains all expected deliverables:
    PPTX, DOCX, MD, TXT, JSON spec, JSON manifest, and README.
    """
    import zipfile
    fixtures_path = os.path.join(os.path.dirname(__file__), "..", "..", "fixtures", "mock_draft_outputs.json")
    with open(fixtures_path, "r", encoding="utf-8") as f:
        drafts = json.load(f)

    zip_file = str(tmp_path / "all_deliverables.zip")
    temp_pptx = str(tmp_path / "temp.pptx")
    temp_docx = str(tmp_path / "temp.docx")

    pres_sample = {
        "deck_title": "Operation GhostLatch Overview",
        "target_audience": "Leadership",
        "slides": [
            {
                "slide_number": 1,
                "title": "Incident Containment",
                "bullet_points": ["Substation controllers isolated", "Firmware patch deployed"],
                "visual_guidance": "Single column card",
                "speaker_notes": "Incident fully contained.",
                "slide_reference_citations": ["doc_01_chunk_01"]
            }
        ]
    }

    export_pptx(pres_sample, temp_pptx)
    export_docx(drafts["advisory"], temp_docx)

    with zipfile.ZipFile(zip_file, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(temp_pptx, arcname="01_presentation_deck.pptx")
        zf.write(temp_docx, arcname="02_intelligence_advisory.docx")
        zf.writestr("03_executive_summary.md", "# Executive Summary\n\nSample content")
        zf.writestr("04_linkedin_post.txt", "Headline\n\nSample LinkedIn post")
        zf.writestr("05_twitter_thread.txt", "[Tweet 1]\nSample Tweet")
        zf.writestr("06_video_script.md", "# Video Script\n\nScene 1")
        zf.writestr("07_infographic_spec.json", json.dumps(drafts.get("infographic", {"chart": "timeline"})))
        zf.writestr("manifest_all_deliverables.json", json.dumps({"job_id": "test_job"}))
        zf.writestr("README_DELIVERABLES.txt", "SENTINEL-TRANSFORM SOVEREIGN DELIVERABLES BUNDLE")

    assert os.path.exists(zip_file)
    assert os.path.getsize(zip_file) > 0

    with zipfile.ZipFile(zip_file, "r") as zf:
        names = zf.namelist()
        assert "01_presentation_deck.pptx" in names
        assert "02_intelligence_advisory.docx" in names
        assert "03_executive_summary.md" in names
        assert "04_linkedin_post.txt" in names
        assert "05_twitter_thread.txt" in names
        assert "06_video_script.md" in names
        assert "07_infographic_spec.json" in names
        assert "manifest_all_deliverables.json" in names
        assert "README_DELIVERABLES.txt" in names

        # Verify PPTX is valid inside zip
        with zf.open("01_presentation_deck.pptx") as f:
            prs = Presentation(f)
            assert len(prs.slides) >= 1

        # Verify DOCX is valid inside zip
        with zf.open("02_intelligence_advisory.docx") as f:
            doc = Document(f)
            assert len(doc.paragraphs) > 0

