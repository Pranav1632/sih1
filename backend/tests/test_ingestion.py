"""
backend/tests/test_ingestion.py - Ingestion Subsystem End-to-End Tests.
Owned by: Ingestion Engineer.

Validates:
- Task 1: SEI Store SQLite schema & zero-loss roundtrip against fixtures/mock_source_chunks.json.
- Task 2: Multimodal Parsers (PDF, DOCX, Image, Audio/Video, Text).
- Task 3: Coordinate-aware chunker & Source Governance (1.0 vs 0.5 weights + numerical conflict detection).
- Task 4: run_ingestion_and_normalization LangGraph AgentState node.
"""

import json
from pathlib import Path
import tempfile
import pytest

from backend.ingestion.chunker import chunk_text, extract_entities_from_text
from backend.ingestion.normalizer import run_ingestion_and_normalization
from backend.ingestion.parsers import (
    parse_audio_video,
    parse_docx,
    parse_file,
    parse_image,
    parse_pdf,
    parse_text,
)
from backend.ingestion.sei_store import SEIStore
from backend.ingestion.source_governance import apply_source_governance, detect_conflicts


@pytest.fixture
def mock_chunks():
    """Loads frozen fixture from fixtures/mock_source_chunks.json."""
    fixture_path = Path("fixtures/mock_source_chunks.json")
    with open(fixture_path, "r", encoding="utf-8") as f:
        return json.load(f)


# =====================================================================
# TASK 1: SQLite Source Evidence Index (SEI Store) Tests
# =====================================================================

def test_sei_store_initialization():
    """Verifies that an in-memory or file-backed SEI store initializes tables correctly."""
    store = SEIStore(db_path=":memory:")
    assert store.count_chunks() == 0


def test_sei_store_roundtrip_fixture_zero_field_loss(mock_chunks):
    """
    Acceptance Criteria (Task 1):
    fixture data round-trips through insert -> query with no field loss.
    """
    store = SEIStore(db_path=":memory:")

    inserted_count = store.insert_chunks(mock_chunks)
    assert inserted_count == len(mock_chunks)
    assert store.count_chunks() == len(mock_chunks)

    # Verify each chunk roundtrips with 100% field fidelity
    for original in mock_chunks:
        retrieved = store.get_chunk(original["chunk_id"])
        assert retrieved is not None, f"Chunk {original['chunk_id']} not found in SEI store"

        # Check all locked schema fields
        for field in [
            "chunk_id",
            "doc_id",
            "source_name",
            "source_role",
            "page_number",
            "timestamp_start",
            "timestamp_end",
            "char_start",
            "char_end",
            "text",
        ]:
            assert retrieved[field] == original[field], (
                f"Field mismatch on '{field}' for {original['chunk_id']}: "
                f"expected {original[field]}, got {retrieved[field]}"
            )

        # Compare extracted_entities (list of dicts)
        assert retrieved["extracted_entities"] == original["extracted_entities"]

        # Ensure authority weight was assigned appropriately
        if original["source_role"] == "PRIMARY":
            assert retrieved["authority_weight"] == 1.0
        else:
            assert retrieved["authority_weight"] == 0.5


def test_sei_store_queries(mock_chunks):
    """Verifies querying chunks by doc_id and by source_role."""
    store = SEIStore(db_path=":memory:")
    store.insert_chunks(mock_chunks)

    # Query by doc_id
    doc_01_chunks = store.get_chunks_by_doc("doc_01")
    assert len(doc_01_chunks) == 2
    assert all(c["doc_id"] == "doc_01" for c in doc_01_chunks)

    doc_02_chunks = store.get_chunks_by_doc("doc_02")
    assert len(doc_02_chunks) == 1
    assert doc_02_chunks[0]["doc_id"] == "doc_02"

    # Query by source_role
    primary_chunks = store.get_chunks_by_role("PRIMARY")
    assert len(primary_chunks) == 2
    assert all(c["source_role"] == "PRIMARY" for c in primary_chunks)

    supporting_chunks = store.get_chunks_by_role("SUPPORTING")
    assert len(supporting_chunks) == 1
    assert supporting_chunks[0]["source_role"] == "SUPPORTING"

    # Clear store
    store.clear_store()
    assert store.count_chunks() == 0


# =====================================================================
# TASK 3: Coordinate-Aware Chunker Tests
# =====================================================================

def test_chunker_small_text_single_chunk():
    """Text within max_chunk_size produces exactly one coordinate-indexed chunk."""
    text = "On 03 September 2026, the Cyber Resilience Unit detected unauthorized access attempts."
    coords = {
        "doc_id": "test_doc",
        "source_name": "incident.pdf",
        "source_role": "PRIMARY",
        "page_number": 1,
        "base_char_offset": 100,
    }
    chunks = chunk_text(text, coordinates=coords, max_chunk_size=500)

    assert len(chunks) == 1
    c = chunks[0]
    assert c["chunk_id"] == "test_doc_chunk_01"
    assert c["doc_id"] == "test_doc"
    assert c["source_role"] == "PRIMARY"
    assert c["page_number"] == 1
    assert c["char_start"] == 100
    assert c["char_end"] == 100 + len(text)
    assert c["text"] == text


def test_chunker_long_text_multi_chunk_boundaries():
    """Long text splits gracefully on sentence boundaries with accurate offsets."""
    sentences = [
        "First sentence detailing the cyber security incident on the regional electric grid.",
        "Second sentence describing the firmware exploit found on substation control nodes.",
        "Third sentence noting the remediation actions taken by the emergency response unit.",
        "Fourth sentence confirming that no active data exfiltration was observed.",
    ]
    long_text = " ".join(sentences)
    coords = {
        "doc_id": "doc_long",
        "source_name": "report.pdf",
        "base_char_offset": 0,
    }
    chunks = chunk_text(long_text, coordinates=coords, max_chunk_size=120, overlap=20)

    assert len(chunks) > 1
    for i, c in enumerate(chunks):
        assert c["chunk_id"] == f"doc_long_chunk_{i+1:02d}"
        assert c["char_start"] >= 0
        assert c["char_end"] > c["char_start"]
        assert len(c["text"]) > 0


def test_chunker_extracts_entities():
    """Verifies that entity extraction identifies named entities in chunks."""
    text = "The Cyber Resilience Unit investigated the intrusion into the substation control software."
    entities = extract_entities_from_text(text)
    assert isinstance(entities, list)
    assert len(entities) > 0
    entity_texts = [e["text"] for e in entities]
    assert any("Cyber" in t or "Resilience" in t for t in entity_texts)


# =====================================================================
# TASK 3: Source Governance & Conflict Detection Tests
# =====================================================================

def test_source_governance_authority_weighting(mock_chunks):
    """
    Acceptance Criteria (Task 3):
    1.0/0.5 authority weighting assigned based on operator-designated primary_doc_id.
    """
    weighted_chunks, conflicts = apply_source_governance(mock_chunks, primary_doc_id="doc_01")

    for c in weighted_chunks:
        if c["doc_id"] == "doc_01":
            assert c["source_role"] == "PRIMARY"
            assert c["authority_weight"] == 1.0
        else:
            assert c["source_role"] == "SUPPORTING"
            assert c["authority_weight"] == 0.5


def test_source_governance_numerical_conflict_detection(mock_chunks):
    """
    Acceptance Criteria (Task 3):
    Test case with contradicting primary/supporting numbers (e.g. '2 endpoints' vs '14 endpoints')
    correctly flags the conflict and keeps primary's number authoritative.
    """
    weighted_chunks, conflicts = apply_source_governance(mock_chunks, primary_doc_id="doc_01")

    assert len(conflicts) >= 1, "Expected numerical conflict on 'endpoints' to be detected."
    conflict = conflicts[0]

    assert conflict["conflict_type"] == "NUMERICAL_DISCREPANCY"
    assert "endpoint" in conflict["subject"]
    assert conflict["primary_value"] == 2
    assert conflict["supporting_value"] == 14
    assert conflict["primary_doc_id"] == "doc_01"
    assert conflict["supporting_doc_id"] == "doc_02"
    assert conflict["resolution"] == "PRIMARY_AUTHORITATIVE"
    assert conflict["authority_weight_applied"] == 1.0


# =====================================================================
# TASK 2: Multimodal Parsers Tests
# =====================================================================

def test_parse_text_file(tmp_path):
    """Tests parsing a plain text file."""
    text_file = tmp_path / "intel_brief.txt"
    content = "Advisory: Critical vulnerability detected in SCADA management gateway."
    text_file.write_text(content, encoding="utf-8")

    chunks = parse_text(str(text_file), doc_id="doc_txt", source_role="PRIMARY")
    assert len(chunks) == 1
    assert chunks[0]["doc_id"] == "doc_txt"
    assert chunks[0]["text"] == content
    assert chunks[0]["source_role"] == "PRIMARY"


def test_parse_pdf_file(tmp_path):
    """Tests creating and parsing a real PDF with PyMuPDF."""
    import pymupdf

    pdf_file = tmp_path / "sample_report.pdf"
    doc = pymupdf.open()

    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Incident Report Page 1: Substations 1 and 2 affected.")

    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Incident Report Page 2: Emergency firmware patch deployed successfully.")

    doc.save(str(pdf_file))
    doc.close()

    chunks = parse_pdf(str(pdf_file), doc_id="doc_pdf", source_role="PRIMARY")
    assert len(chunks) >= 2
    assert chunks[0]["page_number"] == 1
    assert chunks[1]["page_number"] == 2
    assert "Substations" in chunks[0]["text"]
    assert "firmware" in chunks[1]["text"]


def test_parse_docx_file(tmp_path):
    """Tests creating and parsing a real DOCX with python-docx."""
    import docx

    docx_file = tmp_path / "sample_threat.docx"
    doc = docx.Document()
    doc.add_heading("Threat Assessment: Operation GhostLatch", level=1)
    doc.add_paragraph("Preliminary analysis suggests firmware injection via exposed management port.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Indicator"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "C2 IP"
    table.cell(1, 1).text = "192.0.2.14"

    doc.save(str(docx_file))

    chunks = parse_docx(str(docx_file), doc_id="doc_docx", source_role="PRIMARY")
    assert len(chunks) >= 1
    combined_text = " ".join(c["text"] for c in chunks)
    assert "Operation GhostLatch" in combined_text
    assert "firmware injection" in combined_text
    assert "192.0.2.14" in combined_text


def test_parse_image_file(tmp_path):
    """Tests parsing an image file with Pillow contrast enhancement + OCR wrapper."""
    from PIL import Image, ImageDraw

    img_file = tmp_path / "threat_scan.png"
    img = Image.new("RGB", (300, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((10, 40), "NTRO RESTRICTED INTELLIGENCE", fill=(0, 0, 0))
    img.save(str(img_file))

    chunks = parse_image(str(img_file), doc_id="doc_img", source_role="SUPPORTING")
    assert len(chunks) == 1
    assert chunks[0]["doc_id"] == "doc_img"
    assert chunks[0]["source_role"] == "SUPPORTING"
    assert chunks[0]["page_number"] == 1
    assert len(chunks[0]["text"]) > 0


def test_parse_audio_video_interface(tmp_path):
    """Tests audio/video parser output contract, mock model injection, and fallback."""
    dummy_audio = tmp_path / "intercept.mp3"
    dummy_audio.write_bytes(b"dummy audio bytes")

    # 1. Test with injected Whisper model (simulating faster-whisper transcription)
    class MockSegment:
        def __init__(self, start, end, text):
            self.start = start
            self.end = end
            self.text = text

    class MockWhisper:
        def transcribe(self, file_path, beam_size=5):
            return [
                MockSegment(0.0, 14.2, "Intercepted transmission: Unauthorized remote access detected."),
                MockSegment(14.2, 28.5, "Response team deployed emergency mitigation firmware.")
            ], None

    chunks_with_model = parse_audio_video(
        str(dummy_audio),
        doc_id="doc_audio",
        source_role="SUPPORTING",
        model_instance=MockWhisper(),
    )
    assert len(chunks_with_model) == 2
    c1, c2 = chunks_with_model[0], chunks_with_model[1]
    assert c1["chunk_id"] == "doc_audio_chunk_01"
    assert c1["timestamp_start"] == 0.0
    assert c1["timestamp_end"] == 14.2
    assert "Unauthorized remote access" in c1["text"]
    assert c2["timestamp_start"] == 14.2
    assert c2["timestamp_end"] == 28.5

    # 2. Test fallback mode (ensures air-gapped operation when weights are absent)
    chunks_fallback = parse_audio_video(str(dummy_audio), doc_id="doc_audio_fallback", source_role="SUPPORTING")
    assert len(chunks_fallback) >= 1
    chunk = chunks_fallback[0]
    assert chunk["doc_id"] == "doc_audio_fallback"
    assert chunk["timestamp_start"] is not None
    assert chunk["timestamp_end"] is not None
    assert "text" in chunk


def test_parse_file_auto_dispatcher(tmp_path):
    """Tests parse_file dispatching correctly according to extension."""
    txt = tmp_path / "notes.txt"
    txt.write_text("Plaintext dispatcher test", encoding="utf-8")
    chunks = parse_file(str(txt), doc_id="doc_disp")
    assert len(chunks) == 1
    assert "Plaintext dispatcher" in chunks[0]["text"]


# =====================================================================
# TASK 4: LangGraph AgentState Node Tests
# =====================================================================

def test_run_ingestion_and_normalization_with_uploaded_files(tmp_path):
    """
    Acceptance Criteria (Task 4):
    run_ingestion_and_normalization matches stub signature, ingests uploaded_files,
    writes source_chunks, primary_doc_id, extracted_entities, and merged_context conflicts.
    """
    f1 = tmp_path / "primary_incident.txt"
    f1.write_text("Primary: Exactly 2 endpoints were confirmed infected across the power grid.", encoding="utf-8")

    f2 = tmp_path / "secondary_news.txt"
    f2.write_text("Secondary: News sources claim 14 endpoints were affected in the incident.", encoding="utf-8")

    state = {
        "job_id": "test_job_101",
        "uploaded_files": [
            {"path": str(f1), "doc_id": "doc_01", "source_name": "primary_incident.txt", "source_role": "PRIMARY"},
            {"path": str(f2), "doc_id": "doc_02", "source_name": "secondary_news.txt", "source_role": "SUPPORTING"},
        ],
        "primary_doc_id": "doc_01",
        "source_chunks": [],
    }

    test_store = SEIStore(db_path=":memory:")
    result = run_ingestion_and_normalization(state, sei_store=test_store)

    assert "source_chunks" in result
    assert "primary_doc_id" in result
    assert "uploaded_files" in result
    assert "extracted_entities" in result
    assert "merged_context" in result

    assert result["primary_doc_id"] == "doc_01"
    assert len(result["source_chunks"]) == 2

    # Check SEI store was populated
    assert test_store.count_chunks() == 2

    # Check conflict was detected
    conflicts = result["merged_context"].get("conflicts", [])
    assert len(conflicts) >= 1
    assert conflicts[0]["primary_value"] == 2
    assert conflicts[0]["supporting_value"] == 14
    assert conflicts[0]["resolution"] == "PRIMARY_AUTHORITATIVE"


def test_run_ingestion_and_normalization_with_in_memory_fixture(mock_chunks):
    """Tests running the node with pre-existing or in-memory chunks."""
    state = {
        "job_id": "test_job_102",
        "uploaded_files": [],
        "primary_doc_id": "doc_01",
        "source_chunks": mock_chunks,
    }

    test_store = SEIStore(db_path=":memory:")
    result = run_ingestion_and_normalization(state, sei_store=test_store)

    assert len(result["source_chunks"]) == len(mock_chunks)
    assert test_store.count_chunks() == len(mock_chunks)
    assert len(result["merged_context"]["conflicts"]) >= 1
