import io
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from docx import Document

from backend.main import app

client = TestClient(app)


def test_health_check():
    """Verify backend health endpoint."""
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "healthy", "service": "sentinel-transform-backend"}


def test_ingest_endpoint():
    """
    POST /api/ingest
    Validates multipart upload and returns {job_id, source_chunks[]}.
    """
    dummy_pdf = io.BytesIO(b"%PDF-1.4 Mock classified report content for ingestion")
    dummy_pdf.name = "Operation_GhostLatch_Report.pdf"

    response = client.post(
        "/api/ingest",
        files={"files": (dummy_pdf.name, dummy_pdf, "application/pdf")},
        data={"source_role": "PRIMARY"}
    )

    assert response.status_code == 200
    data = response.json()
    assert "job_id" in data
    assert "source_chunks" in data
    assert isinstance(data["source_chunks"], list)
    assert len(data["source_chunks"]) > 0


def test_generate_and_unlocked_export_flow():
    """
    Validates the normal end-to-end flow when Hard Gate is not triggered:
    ingest -> generate -> status -> export (.pptx and .docx).
    """
    # 1. Ingest
    pdf_bytes = io.BytesIO(b"%PDF-1.4 Unlocked flow test")
    ingest_resp = client.post(
        "/api/ingest",
        files={"files": ("test_doc.pdf", pdf_bytes, "application/pdf")},
        data={"source_role": "PRIMARY"}
    )
    job_id = ingest_resp.json()["job_id"]

    # 2. Generate
    gen_resp = client.post(
        "/api/generate",
        json={
            "job_id": job_id,
            "parameters": {"tone": "forensic", "detail": "high"},
            "requested_formats": ["presentation", "advisory"]
        }
    )
    assert gen_resp.status_code == 200
    assert gen_resp.json()["job_id"] == job_id
    assert gen_resp.json()["status"] == "completed"

    # 3. Status
    status_resp = client.get(f"/api/status/{job_id}")
    assert status_resp.status_code == 200
    assert status_resp.json()["status"] == "completed"
    assert status_resp.json()["hard_gate_triggered"] is False

    # 4. Export PPTX
    pptx_resp = client.get(f"/api/export/pptx/{job_id}")
    assert pptx_resp.status_code == 200
    assert "presentationml" in pptx_resp.headers["content-type"]
    prs = Presentation(io.BytesIO(pptx_resp.content))
    assert len(prs.slides) >= 1

    # 5. Export DOCX
    docx_resp = client.get(f"/api/export/docx/{job_id}")
    assert docx_resp.status_code == 200
    assert "wordprocessingml" in docx_resp.headers["content-type"]
    doc = Document(io.BytesIO(docx_resp.content))
    assert len(doc.paragraphs) > 0


def test_hard_gate_export_locking_and_resume():
    """
    Task 4 & Hard Gate Operational Protocol Rule 1:
    - When hard_gate_triggered == True and human_approved == False:
      export endpoint MUST return HTTP 423 (Locked).
    - When human_approved == True via /api/review/confirm:
      export endpoint succeeds with HTTP 200.
    """
    # 1. Ingest
    pdf_bytes = io.BytesIO(b"%PDF-1.4 Hard Gate test")
    ingest_resp = client.post(
        "/api/ingest",
        files={"files": ("mismatch_report.pdf", pdf_bytes, "application/pdf")},
        data={"source_role": "PRIMARY"}
    )
    job_id = ingest_resp.json()["job_id"]

    # 2. Generate with simulated Hard Gate discrepancy
    gen_resp = client.post(
        "/api/generate",
        json={
            "job_id": job_id,
            "parameters": {"simulate_hard_gate": True},
            "requested_formats": ["presentation", "advisory"]
        }
    )
    assert gen_resp.status_code == 200
    assert gen_resp.json()["status"] == "gate_paused"

    # 3. Status shows gate triggered
    status_resp = client.get(f"/api/status/{job_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["status"] == "gate_paused"
    assert status_data["hard_gate_triggered"] is True
    assert len(status_data["entity_discrepancies"]) > 0

    # 4. Attempt export before approval -> MUST return HTTP 423 (Locked)
    lock_resp_pptx = client.get(f"/api/export/pptx/{job_id}")
    assert lock_resp_pptx.status_code == 423
    assert "Export locked: Hard Gate triggered" in lock_resp_pptx.json()["detail"]

    lock_resp_docx = client.get(f"/api/export/docx/{job_id}")
    assert lock_resp_docx.status_code == 423
    assert "Export locked: Hard Gate triggered" in lock_resp_docx.json()["detail"]

    # 5. Operator reviews and confirms in modal: human_approved=True
    confirm_resp = client.post(
        "/api/review/confirm",
        json={
            "job_id": job_id,
            "human_approved": True,
            "human_corrections": {
                "Directorate of Grid Power Resilience": "Directorate of Power Grid Resilience"
            }
        }
    )
    assert confirm_resp.status_code == 200
    assert confirm_resp.json()["status"] == "resumed"

    # 6. Status now shows unblocked
    status_resp2 = client.get(f"/api/status/{job_id}")
    assert status_resp2.status_code == 200
    assert status_resp2.json()["hard_gate_triggered"] is False
    assert status_resp2.json()["status"] == "completed"

    # 7. Export now succeeds with HTTP 200
    unlocked_pptx = client.get(f"/api/export/pptx/{job_id}")
    assert unlocked_pptx.status_code == 200
    assert len(unlocked_pptx.content) > 0

    unlocked_docx = client.get(f"/api/export/docx/{job_id}")
    assert unlocked_docx.status_code == 200
    assert len(unlocked_docx.content) > 0
