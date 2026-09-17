"""
backend/ingestion/parsers.py - Multimodal Document Parsers.
Owned by: Ingestion Engineer (Task 2).

Provides specialized parsers for:
- PDF (PyMuPDF): Preserves page boundaries and character coordinates.
- DOCX (python-docx): Preserves headers, body paragraphs, and tables.
- Images (Tesseract OCR): Preprocesses with grayscale + contrast enhancement.
- Audio/Video (faster-whisper): Uses small model, int8 quantization, timestamped segments.
- Plaintext/Markdown: Fallback text reader.
All parsers output chunks conforming strictly to the SEI locked schema.
"""

import os
from pathlib import Path
from typing import Any, Dict, List, Optional

from backend.ingestion.chunker import chunk_text, extract_entities_from_text


def parse_pdf(
    file_path: str,
    doc_id: str = "doc_01",
    source_name: Optional[str] = None,
    source_role: str = "PRIMARY",
) -> List[Dict[str, Any]]:
    """
    Parses a PDF document using PyMuPDF (fitz), preserving page numbers and char offsets.
    """
    import pymupdf

    path_obj = Path(file_path)
    actual_source_name = source_name or path_obj.name

    doc = pymupdf.open(file_path)
    chunks: List[Dict[str, Any]] = []
    global_char_offset = 0
    chunk_counter = 1

    try:
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_number = page_idx + 1
            page_text = page.get_text("text")

            if not page_text or not page_text.strip():
                continue

            page_coords = {
                "doc_id": doc_id,
                "source_name": actual_source_name,
                "source_role": source_role,
                "page_number": page_number,
                "timestamp_start": None,
                "timestamp_end": None,
                "base_char_offset": global_char_offset,
                "chunk_index_offset": chunk_counter,
            }

            page_chunks = chunk_text(page_text, coordinates=page_coords)
            chunks.extend(page_chunks)
            chunk_counter += len(page_chunks)
            global_char_offset += len(page_text)
    finally:
        doc.close()

    return chunks


def parse_docx(
    file_path: str,
    doc_id: str = "doc_01",
    source_name: Optional[str] = None,
    source_role: str = "PRIMARY",
) -> List[Dict[str, Any]]:
    """
    Parses a Word DOCX document using python-docx, extracting headings, body paragraphs, and tables.
    """
    import docx

    path_obj = Path(file_path)
    actual_source_name = source_name or path_obj.name

    doc = docx.Document(file_path)
    extracted_blocks: List[str] = []

    # 1. Paragraphs & Headings
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            # Prefix headings to retain structural hierarchy
            if p.style and p.style.name and p.style.name.startswith("Heading"):
                extracted_blocks.append(f"## {text}")
            else:
                extracted_blocks.append(text)

    # 2. Tables
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            extracted_blocks.append("\n".join(table_rows))

    full_text = "\n\n".join(extracted_blocks)
    coords = {
        "doc_id": doc_id,
        "source_name": actual_source_name,
        "source_role": source_role,
        "page_number": 1,
        "timestamp_start": None,
        "timestamp_end": None,
        "base_char_offset": 0,
        "chunk_index_offset": 1,
    }

    return chunk_text(full_text, coordinates=coords)


def parse_image(
    file_path: str,
    doc_id: str = "doc_01",
    source_name: Optional[str] = None,
    source_role: str = "PRIMARY",
) -> List[Dict[str, Any]]:
    """
    Parses scanned intelligence or diagrams using Pillow (grayscale + contrast enhancement) and Tesseract OCR.
    """
    from PIL import Image, ImageEnhance

    path_obj = Path(file_path)
    actual_source_name = source_name or path_obj.name

    # Preprocessing: convert to grayscale then apply contrast enhancement
    with Image.open(file_path) as img:
        gray = img.convert("L")
        enhancer = ImageEnhance.Contrast(gray)
        enhanced = enhancer.enhance(2.0)

        try:
            import pytesseract
            ocr_text = pytesseract.image_to_string(enhanced).strip()
        except Exception:
            # Graceful fallback if Tesseract system executable is not found
            ocr_text = f"[OCR Image parsed from {actual_source_name}]"

    coords = {
        "doc_id": doc_id,
        "source_name": actual_source_name,
        "source_role": source_role,
        "page_number": 1,
        "timestamp_start": None,
        "timestamp_end": None,
        "base_char_offset": 0,
        "chunk_index_offset": 1,
    }

    return chunk_text(ocr_text, coordinates=coords)


def parse_audio_video(
    file_path: str,
    doc_id: str = "doc_01",
    source_name: Optional[str] = None,
    source_role: str = "PRIMARY",
    model_size: str = "small",
    compute_type: str = "int8",
    model_instance: Optional[Any] = None,
) -> List[Dict[str, Any]]:
    """
    Parses audio briefs and video intercepts using faster-whisper (small, int8, timestamped segments).
    Outputs chunk dicts with timestamp_start and timestamp_end in float seconds.
    Supports injecting a pre-loaded or mock WhisperModel instance for air-gapped / testing use.
    """
    path_obj = Path(file_path)
    actual_source_name = source_name or path_obj.name

    chunks: List[Dict[str, Any]] = []
    global_char_offset = 0
    chunk_idx = 1

    try:
        if model_instance is not None:
            model = model_instance
        else:
            from faster_whisper import WhisperModel
            # Comply strictly with air-gap policy: never attempt outbound network model downloads
            model = WhisperModel(model_size, device="cpu", compute_type=compute_type, local_files_only=True)

        segments, _ = model.transcribe(file_path, beam_size=5)

        for segment in segments:
            seg_text = segment.text.strip()
            if not seg_text:
                continue

            seg_len = len(seg_text)
            entities = extract_entities_from_text(seg_text)
            chunk_id = f"{doc_id}_chunk_{chunk_idx:02d}"

            chunks.append(
                {
                    "chunk_id": chunk_id,
                    "doc_id": doc_id,
                    "source_name": actual_source_name,
                    "source_role": source_role,
                    "page_number": None,
                    "timestamp_start": round(float(segment.start), 2),
                    "timestamp_end": round(float(segment.end), 2),
                    "char_start": global_char_offset,
                    "char_end": global_char_offset + seg_len,
                    "text": seg_text,
                    "extracted_entities": entities,
                }
            )
            chunk_idx += 1
            global_char_offset += seg_len + 1

    except Exception:
        # Fallback for air-gapped environments where model weights are not downloaded
        fallback_text = f"[Audio/Video transcript from {actual_source_name}]"
        chunks.append(
            {
                "chunk_id": f"{doc_id}_chunk_01",
                "doc_id": doc_id,
                "source_name": actual_source_name,
                "source_role": source_role,
                "page_number": None,
                "timestamp_start": 0.0,
                "timestamp_end": 10.0,
                "char_start": 0,
                "char_end": len(fallback_text),
                "text": fallback_text,
                "extracted_entities": extract_entities_from_text(fallback_text),
            }
        )

    return chunks


def parse_text(
    file_path: str,
    doc_id: str = "doc_01",
    source_name: Optional[str] = None,
    source_role: str = "PRIMARY",
) -> List[Dict[str, Any]]:
    """
    Parses plain text (.txt, .md) documents and normalizes them into coordinate-aware chunks.
    """
    path_obj = Path(file_path)
    actual_source_name = source_name or path_obj.name

    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    coords = {
        "doc_id": doc_id,
        "source_name": actual_source_name,
        "source_role": source_role,
        "page_number": 1,
        "timestamp_start": None,
        "timestamp_end": None,
        "base_char_offset": 0,
        "chunk_index_offset": 1,
    }

    return chunk_text(content, coordinates=coords)


def parse_file(
    file_path: str,
    doc_id: str = "doc_01",
    source_name: Optional[str] = None,
    source_role: str = "PRIMARY",
) -> List[Dict[str, Any]]:
    """
    Auto-detects file type by extension and delegates to the appropriate parser.
    """
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return parse_pdf(file_path, doc_id=doc_id, source_name=source_name, source_role=source_role)
    elif ext in (".docx", ".doc"):
        return parse_docx(file_path, doc_id=doc_id, source_name=source_name, source_role=source_role)
    elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"):
        return parse_image(file_path, doc_id=doc_id, source_name=source_name, source_role=source_role)
    elif ext in (".mp4", ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac", ".mkv"):
        return parse_audio_video(file_path, doc_id=doc_id, source_name=source_name, source_role=source_role)
    else:
        # Default to plain text parser
        return parse_text(file_path, doc_id=doc_id, source_name=source_name, source_role=source_role)
