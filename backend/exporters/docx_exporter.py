import os
from typing import Union, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from backend.models.formats.advisory_schema import AdvisorySchema


def _set_cell_background(cell, hex_color: str):
    """Sets the background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def export_docx(advisory: Union[AdvisorySchema, Dict[str, Any]], output_path: str) -> str:
    """
    Exports an AdvisorySchema (or compatible dict) to a formal Intelligence Advisory (.docx) file.
    Includes institutional header, severity color banner, IOC table, and footnote citations.
    """
    if isinstance(advisory, dict):
        adv = AdvisorySchema(**advisory)
    else:
        adv = advisory

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    doc = Document()

    # Configure Margins (0.75 in for sovereign intelligence brief layout)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Institutional Header
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = Inches(4.5)
    cell_right.width = Inches(2.5)

    p_org = cell_left.paragraphs[0]
    r_org = p_org.add_run("NATIONAL TECHNICAL RESEARCH ORGANISATION (NTRO)")
    r_org.bold = True
    r_org.font.size = Pt(10)
    r_org.font.color.rgb = RGBColor(30, 41, 59)
    p_unit = cell_left.add_paragraph()
    r_unit = p_unit.add_run("CYBER DEFENSE & CRITICAL INFRASTRUCTURE RESILIENCE DIRECTORY")
    r_unit.font.size = Pt(8.5)
    r_unit.font.color.rgb = RGBColor(100, 116, 139)

    p_adv_id = cell_right.paragraphs[0]
    p_adv_id.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_adv_id = p_adv_id.add_run(f"REF: {adv.advisory_id}")
    r_adv_id.bold = True
    r_adv_id.font.size = Pt(11)
    r_adv_id.font.color.rgb = RGBColor(15, 23, 42)
    p_class = cell_right.add_paragraph()
    p_class.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_class = p_class.add_run("AUTHENTICATED // SOVEREIGN TRANSFORM")
    r_class.font.size = Pt(8.5)
    r_class.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()  # Spacing

    # 2. Document Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(adv.title)
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # 3. Severity Color Banner
    severity_map = {
        "CRITICAL": {"bg": "7F1D1D", "fg": RGBColor(255, 255, 255), "desc": "IMMEDIATE MANDATORY ACTION REQUIRED"},
        "HIGH": {"bg": "B91C1C", "fg": RGBColor(255, 255, 255), "desc": "ELEVATED THREAT - EXPEDITED PATCH / MITIGATION REQUIRED"},
        "MEDIUM": {"bg": "D97706", "fg": RGBColor(255, 255, 255), "desc": "MODERATE RISK - REMEDIATE WITHIN STANDARD CYCLE"},
        "LOW": {"bg": "047857", "fg": RGBColor(255, 255, 255), "desc": "INFORMATIONAL / ROUTINE AUDIT POSTURE"}
    }
    sev_info = severity_map.get(adv.severity_level.upper(), severity_map["HIGH"])

    sev_table = doc.add_table(rows=1, cols=1)
    sev_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sev_cell = sev_table.cell(0, 0)
    sev_cell.width = Inches(7.0)
    _set_cell_background(sev_cell, sev_info["bg"])
    _set_cell_margins(sev_cell, top=140, bottom=140, left=180, right=180)

    p_sev = sev_cell.paragraphs[0]
    p_sev.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sev_tag = p_sev.add_run(f"SEVERITY: {adv.severity_level.upper()}  |  ")
    r_sev_tag.bold = True
    r_sev_tag.font.size = Pt(12)
    r_sev_tag.font.color.rgb = sev_info["fg"]

    r_sev_desc = p_sev.add_run(sev_info["desc"])
    r_sev_desc.font.size = Pt(10)
    r_sev_desc.font.color.rgb = sev_info["fg"]

    doc.add_paragraph()  # Spacing

    # 4. Threat Overview
    h_overview = doc.add_heading(level=1)
    r_h1 = h_overview.add_run("1. Executive Threat Overview")
    r_h1.font.size = Pt(13)
    r_h1.font.color.rgb = RGBColor(15, 23, 42)

    p_overview = doc.add_paragraph()
    p_overview.paragraph_format.line_spacing = 1.15
    p_overview.paragraph_format.space_after = Pt(10)
    r_ov = p_overview.add_run(adv.threat_overview)
    r_ov.font.size = Pt(10.5)

    # 5. Affected Systems
    h_aff = doc.add_heading(level=1)
    r_h2 = h_aff.add_run("2. Targeted Systems & Asset Exposure")
    r_h2.font.size = Pt(13)
    r_h2.font.color.rgb = RGBColor(15, 23, 42)

    if adv.affected_systems:
        for system in adv.affected_systems:
            p_sys = doc.add_paragraph(style="List Bullet")
            p_sys.paragraph_format.space_after = Pt(3)
            r_sys = p_sys.add_run(system)
            r_sys.font.size = Pt(10.5)
    else:
        p_none = doc.add_paragraph()
        p_none.add_run("No specific target subsystems enumerated in evidence.").italic = True

    # 6. Indicators of Compromise (IOC) Table
    h_ioc = doc.add_heading(level=1)
    r_h3 = h_ioc.add_run("3. Indicators of Compromise (IOC)")
    r_h3.font.size = Pt(13)
    r_h3.font.color.rgb = RGBColor(15, 23, 42)

    ioc_table = doc.add_table(rows=1, cols=3)
    ioc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ioc_table.autofit = False

    # Header Row
    hdr_cells = ioc_table.rows[0].cells
    hdr_cells[0].width = Inches(0.8)
    hdr_cells[1].width = Inches(4.5)
    hdr_cells[2].width = Inches(1.7)

    headers = ["#", "Indicator / Artifact / Signature", "Context / Type"]
    for i, title in enumerate(headers):
        _set_cell_background(hdr_cells[i], "1E293B")  # Slate 800
        _set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Populate IOC Rows
    if adv.indicators_of_compromise:
        for idx, ioc in enumerate(adv.indicators_of_compromise, 1):
            row_cells = ioc_table.add_row().cells
            row_cells[0].width = Inches(0.8)
            row_cells[1].width = Inches(4.5)
            row_cells[2].width = Inches(1.7)

            for c in row_cells:
                _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                if idx % 2 == 0:
                    _set_cell_background(c, "F8FAFC")

            p_idx = row_cells[0].paragraphs[0]
            p_idx.add_run(str(idx)).font.size = Pt(9.5)

            p_val = row_cells[1].paragraphs[0]
            r_val = p_val.add_run(ioc)
            r_val.font.name = "Consolas"
            r_val.font.size = Pt(9)

            p_ctx = row_cells[2].paragraphs[0]
            p_ctx.add_run("Observational Evidence").font.size = Pt(9)
    else:
        row_cells = ioc_table.add_row().cells
        for c in row_cells:
            _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        row_cells[0].paragraphs[0].add_run("1").font.size = Pt(9.5)
        row_cells[1].paragraphs[0].add_run("Pending threat discovery disclosure").italic = True
        row_cells[2].paragraphs[0].add_run("Pending").font.size = Pt(9)

    doc.add_paragraph()  # Spacing

    # 7. Recommended Mitigations
    h_mit = doc.add_heading(level=1)
    r_h4 = h_mit.add_run("4. Actionable Mitigations & Directives")
    r_h4.font.size = Pt(13)
    r_h4.font.color.rgb = RGBColor(15, 23, 42)

    if adv.recommended_mitigations:
        for mit in adv.recommended_mitigations:
            p_mit = doc.add_paragraph(style="List Bullet")
            p_mit.paragraph_format.space_after = Pt(3)
            r_m = p_mit.add_run(mit)
            r_m.font.size = Pt(10.5)
    else:
        p_none = doc.add_paragraph()
        p_none.add_run("Maintain standard perimeter defenses.").italic = True

    # 8. Compliance & Governance
    h_gov = doc.add_heading(level=1)
    r_h5 = h_gov.add_run("5. Compliance & Reporting Directive")
    r_h5.font.size = Pt(13)
    r_h5.font.color.rgb = RGBColor(15, 23, 42)

    p_gov = doc.add_paragraph()
    p_gov.paragraph_format.line_spacing = 1.15
    p_gov.paragraph_format.space_after = Pt(10)
    r_g = p_gov.add_run(adv.compliance_and_governance)
    r_g.font.size = Pt(10.5)

    # 9. Forensic Footnote Citations
    doc.add_paragraph()
    p_sep = doc.add_paragraph()
    r_sep = p_sep.add_run("—" * 45)
    r_sep.font.color.rgb = RGBColor(203, 213, 225)

    p_cite = doc.add_paragraph()
    r_ch = p_cite.add_run("EVIDENCE ATTRIBUTION & SOURCE CHUNK CITATIONS\n")
    r_ch.bold = True
    r_ch.font.size = Pt(8.5)
    r_ch.font.color.rgb = RGBColor(71, 85, 105)

    if adv.cited_chunk_ids:
        for cid in adv.cited_chunk_ids:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_after = Pt(1)
            r_c = p_item.add_run(f"• [SEI Index Anchor]: {cid}")
            r_c.font.size = Pt(8.5)
            r_c.font.color.rgb = RGBColor(100, 116, 139)
    else:
        p_none = doc.add_paragraph()
        r_c = p_none.add_run("Primary Context Anchor: Synthesized Source Baseline")
        r_c.font.size = Pt(8.5)
        r_c.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(output_path)
    return output_path
